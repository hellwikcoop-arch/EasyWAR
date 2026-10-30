#!/usr/bin/env python3
"""Synthetic ПС, РВіА, ПВК and УДШВ tables for driving EasyWAR.

The repository only ships empty base templates (src-tauri/resources/base-tables).
This fills their «Зведена» sheet with deterministic, obviously fake rows
(«Тестове N» places, random MGRS strings, TS-/5-digit target numbers), so the
dashboards, mission cards and 3D munition models have something to show.
No real data: everything is generated from a fixed seed.

ПВК: журнал «Таблиця 1» у порожньому шаблоні (аркуші «АА», «ПС», «Об'єкти» мають формули без кешованих значень,
тому вони на дашборді порожні). УДШВ: книга з нуля у формі, яку розпізнає udshv-dashboard-core.js.

Usage: python3 .claude/skills/run-easywar/fixtures.py [out_dir]   (default /tmp/easywar-fixtures)
"""
import datetime
import os
import random
import sys

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', '..'))
BASE = os.path.join(ROOT, 'src-tauri', 'resources', 'base-tables')
MONTHS = ['Січень', 'Лютий', 'Березень', 'Квітень', 'Травень', 'Червень',
          'Липень', 'Серпень', 'Вересень', 'Жовтень', 'Листопад', 'Грудень']
UNITS = ['7 ОДШБр', '15 ОАБр', '24 ОМБр', '33 ОМБр', '47 ОМБр', '58 ОМБр', '11 ОМБр']
TARGETS = ['КСП', 'ПУ БПЛА', 'Укриття', 'Переправа', 'Накопичення о/с / ШГ', 'ВОП/РОП',
           "Вузол зв'язку", 'БК/склад', 'Вогнева позиція']
RESULTS = ['ЗНИЩЕНО'] * 5 + ['ПОШКОДЖЕНО'] * 3 + ['ПОДАВЛЕНО', 'НЕ УРАЖЕНО', 'НЕ УРАЖЕНО', 'НЕ ВІДОМО']
AIR_REASONS = ['ППО', 'РЕБ', 'Відмова АЗУ', 'Перехоплено авіацією', 'Потребує дорозвідки']
ROCKET_REASONS = ['ППО', 'РЕБ', 'Перехоплено ППО в польоті', 'Похибка координат', 'Технічна відмова ПУ']
AIRCRAFT = ['Су-25', 'Су-27', 'МіГ-29', 'F-16', 'Mirage 2000']
AIR_MUNITIONS = ['GBU-39', 'GBU-62', 'GBU-64', 'GBU-53/B', 'AASM-250', 'AGM-88 HARM', 'BLU-140',
                 'С-8', 'С-13', 'ОФАБ-250']
LAUNCHERS = {'М142': ['M30A1', 'M30A2', 'M31A1', 'ATACMS'], 'М270': ['M26', 'M30A2', 'M31A1', 'ATACMS']}


def mgrs(rnd):
    # Зона 37U, квадрати 100 км, валідні для непарної зони: стовпці C–F (схід 300–700 км),
    # рядки P–Q (північ ≈ 5300–5500 км), тобто приблизно 48–49,5° пн. ш. Інші пари літер
    # дають точки за тисячі кілометрів, і карта дашборда розтягується.
    return f"37U {rnd.choice('CDEF')}{rnd.choice('PQ')} {rnd.randint(10000, 99999)} {rnd.randint(10000, 99999)}"


def fill(template, out, rows, row_fn, seed):
    rnd = random.Random(seed)
    wb = openpyxl.load_workbook(os.path.join(BASE, template))
    ws = wb['Зведена']
    day = datetime.datetime(2026, 6, 1)
    for n in range(1, rows + 1):
        day += datetime.timedelta(days=rnd.choice([0, 0, 1, 1, 2]))
        result = rnd.choice(RESULTS)
        ws.append(row_fn(n, day, result, rnd))
    wb.save(out)


def air_row(n, day, result, rnd):
    reason = rnd.choice(AIR_REASONS) if result == 'НЕ УРАЖЕНО' else None
    return [n, day, f'{rnd.randint(0, 23):02d}:{rnd.choice(["00", "15", "30", "45"])}', f'TS-{rnd.randint(1000, 9999)}',
            rnd.choice(UNITS), f'Тестове {rnd.randint(1, 14)}', mgrs(rnd), rnd.choice(TARGETS), MONTHS[day.month - 1],
            rnd.choice(AIRCRAFT), rnd.choice(AIR_MUNITIONS), rnd.randint(1, 4), result, reason, None]


def rocket_row(n, day, result, rnd):
    launcher = rnd.choice(list(LAUNCHERS))
    reason = rnd.choice(ROCKET_REASONS) if result == 'НЕ УРАЖЕНО' else None
    return [n, day, f'{rnd.randint(0, 23):02d}:{rnd.choice(["00", "15", "30", "45"])}', str(rnd.randint(10000, 99999)),
            rnd.choice(UNITS), f'Тестове {rnd.randint(1, 14)}', mgrs(rnd), rnd.choice(TARGETS), MONTHS[day.month - 1],
            launcher, rnd.choice(LAUNCHERS[launcher]), rnd.randint(1, 6), result, reason, None,
            rnd.choice(UNITS), f'Тестова ВБр {rnd.randint(1, 3)}']


PVK_TARGETS = ['ЖС ВЗ', 'Арт. Системи', 'ОВТ', 'ППО', 'Рембаза', 'Склад БК']
PVK_FORCES = [('Mi-8', 'ARS-8'), ('Mi-24', 'ARS-8'), ('F-16', 'AGM-88'), ('Су-25', 'С-8 ком')]
PVK_RESULTS = ['Уражено', 'Уражено', 'Не уражено']


def pvk(out, rows=40, seed=7):
    rnd = random.Random(seed)
    wb = openpyxl.load_workbook(os.path.join(BASE, 'pvk-base.xlsx'))
    ws = wb['Таблиця 1']
    day = datetime.datetime(2026, 6, 1)
    for n in range(1, rows + 1):
        day += datetime.timedelta(days=rnd.choice([0, 1, 1, 2]))
        force, ammo = rnd.choice(PVK_FORCES)
        ws.append([n, day, f'{rnd.randint(0, 23):02d}:{rnd.choice(["00", "15", "30", "45"])}', f'TS-{rnd.randint(1000, 9999)}',
                   rnd.choice(PVK_TARGETS), mgrs(rnd), f'Тестове {rnd.randint(1, 14)}', rnd.randint(1, 3), force, ammo,
                   rnd.randint(1, 12), rnd.choice(PVK_RESULTS)])
    wb.save(out)


def udshv(out, days=12, seed=11):
    rnd = random.Random(seed)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'СЕРПЕНЬ-26'
    ws.append(['Втрати о/с та ОВТ противника за угруповання ДШВ'])
    ws.append(['№ з/п', 'День тижня', 'Дата', 'Підрозділ', 'ВгЗ', 'КББ', 'Витрата боєприпасів', None, None, 'Втрати о/с', None,
               'Втрати противника за категоріями', None, 'ЗНИЩЕНО', None, 'ПОШКОДЖЕНО', None, 'ПОДАВЛЕНО', None])
    ws.append([None, None, None, None, None, None, '155мм', '122мм', 'ВСЬОГО', '200', '300', 'артилерія', 'ІФС', 'танк',
               'мотоцикл', 'танк', 'КП, ПУ, КСП', 'танк', 'КП, ПУ, КСП'])
    ws.append([])
    names = ['АГ 148 оабр', 'РУК 148 оабр']
    day = datetime.datetime(2026, 8, 1)
    for _ in range(days):
        for i, name in enumerate(names):
            vals = [rnd.randint(0, 9) for _ in range(5)]
            ws.append([i + 1, 'ДЕНЬ' if i == 0 else None, day if i == 0 else None, name, rnd.randint(0, 9), rnd.randint(0, 4),
                       rnd.randint(0, 40), rnd.randint(0, 30), None, vals[0], vals[1], vals[2], vals[3], vals[4], rnd.randint(0, 3),
                       rnd.randint(0, 3), rnd.randint(0, 2), rnd.randint(0, 2), rnd.randint(0, 2)])
        ws.append([])
        ws.append([])
        day += datetime.timedelta(days=1)
    wb.save(out)


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else '/tmp/easywar-fixtures'
    os.makedirs(out, exist_ok=True)
    fill('air-base.xlsx', os.path.join(out, 'air.xlsx'), 240, air_row, 62)
    fill('rvia-base.xlsx', os.path.join(out, 'rocket.xlsx'), 240, rocket_row, 142)
    pvk(os.path.join(out, 'pvk.xlsx'))
    udshv(os.path.join(out, 'udshv.xlsx'))
    for name in ('air', 'rocket', 'pvk', 'udshv'):
        print(os.path.join(out, name + '.xlsx'))


if __name__ == '__main__':
    main()
