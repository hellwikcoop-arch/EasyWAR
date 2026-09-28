#!/usr/bin/env python3
"""Checks that the ПВК base table carries «Таблиця 1» into ТАБЛ 2-3.

    python3 script/check_pvk_base_formulas.py

Writes synthetic strikes into a copy of src-tauri/resources/base-tables/pvk-base.xlsx,
lets LibreOffice Calc recalculate it (soffice with the Calc component must be
installed) and compares the daily sheets with the values the strikes imply.
Nothing is written into the repository.
"""
import datetime
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent / 'src-tauri' / 'resources' / 'base-tables' / 'pvk-base.xlsx'
DAY = lambda day: datetime.datetime(2026, 10, day)
TIME = lambda hour, minute: datetime.time(hour, minute)
# № з/п, Дата, Час удару, Номер цілі, Характер цілі, Координати цілі, Населений пункт,
# Фактичні приходи, Сили та засоби, Тип АЗУ, Витрата АЗУ, Результат
STRIKES = [
    [1, DAY(1), TIME(10, 15), 'SYN-1', 'ПУ БпЛА', '', 'Тестове', 2, 'Су-25', 'GBU-39', 2, 'Уражено'],
    [2, DAY(1), TIME(12, 0), 'SYN-2', 'Склад БК', '', 'Тестове', 8, 'Mi-24', 'С-8 ком', 10, 'НЕ уражено'],
    [3, DAY(1), TIME(23, 5), 'SYN-3', 'ЖС ВЗ', '', '', 3, 'Су-25', 'C-13', 4, 'Дорозвідується'],
    [4, DAY(2), TIME(6, 30), 'SYN-4', 'ПУ БпЛА', '', '', 6, 'Mi-8', 'C-13', 6, 'Уражено'],
    [5, DAY(4), TIME(9, 0), 'SYN-1', "Вз.зв'язку", '', '', 2, 'F-16', 'AGM-88', 2, 'Уражено'],
    # A carrier the table does not know counts only on «Об'єкти».
    [6, DAY(4), TIME(9, 30), 'SYN-5', 'СП', '', '', 1, 'Ка-52', 'C-13', 3, 'Уражено'],
]
# Non-zero cells per day; everything else on the first four days must be 0.
EXPECTED = {
    'АА': {DAY(1): {'АУ': 1, 'Mi-24': 1, 'С-8 ком': 10, 'НЕ уражено': 1},
           DAY(2): {'АУ': 1, 'Mi-8': 1, 'C-13': 6, 'Уражено': 1}},
    'ПС': {DAY(1): {'АУ': 2, 'Су-25': 2, 'GBU-39': 2, 'C-13': 4, 'Уражено': 1, 'Дорозвідується': 1},
           DAY(4): {'АУ': 1, 'F-16': 1, 'AGM-88': 2, 'Уражено': 1}},
    "Об'єкти": {DAY(1): {'ПУ БпЛА': 1, 'Витрата': 16}, DAY(2): {'ПУ БпЛА': 1, 'Витрата': 6},
                DAY(4): {'СП': 1, "Вз.зв'язку": 1, 'Витрата': 5}},
}


def main():
    work = Path(tempfile.mkdtemp())
    wb = openpyxl.load_workbook(BASE)
    journal = wb['Таблиця 1']
    journal['A2'] = 'ЖОВТЕНЬ'
    for row, strike in enumerate(STRIKES, start=3):
        for column, value in enumerate(strike, start=1):
            if value != '':
                journal.cell(row, column, value)
    wb.save(work / 'pvk.xlsx')
    subprocess.run(['soffice', '--headless', '--norestore', '--convert-to', 'xlsx', '--outdir', str(work / 'out'),
                    str(work / 'pvk.xlsx')], check=True, capture_output=True, env={'HOME': str(work), 'PATH': '/usr/bin:/bin'})
    result = openpyxl.load_workbook(work / 'out' / 'pvk.xlsx', data_only=True)
    problems = []
    for sheet, days in EXPECTED.items():
        ws = result[sheet]
        headers = [str(cell.value).strip() for cell in ws[1]]
        rows = {ws.cell(row, 1).value: row for row in range(2, 6)}
        if list(rows) != [DAY(1), DAY(2), DAY(3), DAY(4)] or ws.cell(6, 1).value not in (None, ''):
            problems.append(f'{sheet}: days {list(rows)}')
        for day, row in rows.items():
            for column, header in enumerate(headers[1:], start=2):
                want = days.get(day, {}).get(header, 0)
                got = ws.cell(row, column).value
                if got != want:
                    problems.append(f'{sheet} {day:%d.%m} «{header}»: {got!r}, очікувалось {want}')
    if problems:
        sys.exit('\n'.join(problems))
    print('ТАБЛ 2-3 збігаються з «Таблицею 1».')


if __name__ == '__main__':
    main()
